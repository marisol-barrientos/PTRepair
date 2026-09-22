#!/usr/bin/env python3
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

ROOT_DIR = Path(
    "/home/marisolbarrientosmoreno/Desktop/BPM_2026_mitigation_actions/ResolvingComplianceViolations/data/declare_constraints_Maggi/"
)


# -----------------------
# Helpers
# -----------------------
def local_name(tag: str) -> str:
    return tag.split("}", 1)[-1] if "}" in tag else tag


def norm(s):
    if s is None:
        return None
    s = str(s).strip()
    return s if s else None


def first(*vals):
    for v in vals:
        v = norm(v)
        if v:
            return v
    return None


def normalize_template(t):
    t = norm(t)
    if not t:
        return None
    t = t.lower().replace("_", " ").strip()
    t = " ".join(t.split())
    return t


def autosize(ws):
    for col in range(1, ws.max_column + 1):
        mx = 0
        letter = get_column_letter(col)
        for row in range(1, ws.max_row + 1):
            v = ws.cell(row=row, column=col).value
            if v is None:
                continue
            mx = max(mx, len(str(v)))
        ws.column_dimensions[letter].width = min(max(12, mx + 2), 80)


def nl(template, a, b):
    t = normalize_template(template) or "unknown"
    a = a or "<?>"
    b = b or "<?>"

    if t == "existence":
        return f"'{a}' must occur at least once."
    if t == "absence":
        return f"'{a}' must not occur."
    if t == "response":
        return f"Whenever '{a}' occurs, '{b}' must occur later."
    if t == "precedence":
        return f"'{b}' may occur only if '{a}' occurred before."
    if t == "succession":
        return f"'{a}' must be followed (eventually) by '{b}', and '{b}' requires a preceding '{a}'."
    if t in {"co existence", "co-existence", "coexistence"}:
        return f"'{a}' and '{b}' must either both occur or both not occur."
    if t in {"not co existence", "not-co-existence", "not coexistence", "notcoexistence"}:
        return f"'{a}' and '{b}' must not both occur."
    if t in {"not succession", "notsuccession"}:
        return f"It is forbidden for '{a}' to be followed later by '{b}'."

    params = [p for p in [a, b] if p and p != "<?>"]
    params_str = ", ".join(f"'{p}'" for p in params)
    return f"{template or 'unknown'}({params_str})"


def extract_template(el):
    t = first(el.get("template"), el.get("type"), el.get("name"),
              el.get("Template"), el.get("Type"))
    if t:
        return t

    for ch in list(el):
        ln = local_name(ch.tag).lower()
        if ln in {"template", "type", "name", "constrainttype"}:
            if norm(ch.text):
                return norm(ch.text)

    return None


def extract_constraintparameters_names(el):
    """
    ConDec Declare XML:
    <constraintparameters>
      <parameter templateparameter="1"><branches><branch name="Activity X"/></branches></parameter>
      ...
    </constraintparameters>
    """
    cp = None
    for ch in list(el):
        if local_name(ch.tag).lower() == "constraintparameters":
            cp = ch
            break
    if cp is None:
        return []

    mapping = {}
    for p in list(cp):
        if local_name(p.tag).lower() != "parameter":
            continue
        tp = p.get("templateparameter")
        if not tp:
            continue

        branch_name = None
        for b in p.iter():
            if local_name(b.tag).lower() == "branch":
                branch_name = first(b.get("name"), b.text)
                if branch_name:
                    break

        if branch_name:
            try:
                mapping[int(tp)] = branch_name
            except ValueError:
                mapping[tp] = branch_name

    keys = sorted(mapping.keys(), key=lambda x: (isinstance(x, str), x))
    return [mapping[k] for k in keys]


# -----------------------
# Core: one file -> one excel
# -----------------------
def declare_to_excel(declare_xml: Path) -> Path:
    root = ET.parse(declare_xml).getroot()

    rows = []
    for el in root.iter():
        ln = local_name(el.tag).lower()

        if ln in {"constraintdefinitions", "constraints", "templates"}:
            continue
        if ln not in {"constraint", "rule", "constraintinstance", "templateconstraint", "declareconstraint"}:
            continue

        template = extract_template(el)
        params = extract_constraintparameters_names(el)

        if not template and not params:
            continue

        a = params[0] if len(params) >= 1 else None
        b = params[1] if len(params) >= 2 else None

        rows.append({
            "template": template,
            "all_params": ", ".join(params) if params else None,
            "natural_language": nl(template, a, b),
        })

    out_xlsx = declare_xml.with_name("declare_summary.xlsx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Constraints"

    headers = ["template", "all_params", "natural_language"]
    ws.append(headers)
    for r in rows:
        ws.append([r.get(h) for h in headers])

    bold = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="center")
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = bold
        cell.alignment = wrap
    ws.freeze_panes = "A2"

    autosize(ws)
    wb.save(out_xlsx)
    return out_xlsx


def main():
    if not ROOT_DIR.exists():
        raise SystemExit(f"Root folder not found: {ROOT_DIR}")

    declare_files = sorted(ROOT_DIR.rglob("declare.xml"))

    if not declare_files:
        print(f"No declare.xml files found under: {ROOT_DIR}")
        return

    print(f"Found {len(declare_files)} declare.xml files.")
    ok = 0
    failed = 0

    for f in declare_files:
        try:
            out = declare_to_excel(f)
            ok += 1
            print(f"[OK] {f} -> {out}")
        except Exception as e:
            failed += 1
            print(f"[FAIL] {f}: {e}")

    print(f"Done. Success: {ok}, Failed: {failed}")


if __name__ == "__main__":
    main()